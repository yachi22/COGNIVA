import urllib.request
import urllib.error
import json
import uuid
import fitz
import docx
import io

def encode_multipart_formdata(fields, files):
    boundary = f"----WebKitFormBoundary{uuid.uuid4().hex}"
    body = bytearray()
    
    for key, value in fields.items():
        body.extend(f"--{boundary}\r\n".encode("utf-8"))
        body.extend(f'Content-Disposition: form-data; name="{key}"\r\n\r\n'.encode("utf-8"))
        body.extend(f"{value}\r\n".encode("utf-8"))
        
    for key, (filename, content, content_type) in files.items():
        body.extend(f"--{boundary}\r\n".encode("utf-8"))
        body.extend(f'Content-Disposition: form-data; name="{key}"; filename="{filename}"\r\n'.encode("utf-8"))
        body.extend(f"Content-Type: {content_type}\r\n\r\n".encode("utf-8"))
        body.extend(content)
        body.extend(b"\r\n")
        
    body.extend(f"--{boundary}--\r\n".encode("utf-8"))
    content_type = f"multipart/form-data; boundary={boundary}"
    return body, content_type


def create_sample_pdf():
    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_text(fitz.Point(50, 100), "Cogniva Intelligence: Page 1 Core Principles and Architecture.", fontsize=12)
    p2 = doc.new_page()
    p2.insert_text(fitz.Point(50, 100), "Cogniva Provenance: Page 2 Citation and Chunking Strategy.", fontsize=12)
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


def create_sample_docx():
    doc = docx.Document()
    doc.add_heading("Enterprise Knowledge Protocol", level=1)
    doc.add_paragraph("This paragraph contains standard body text describing security guidelines.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Segregation"
    table.cell(1, 1).text = "Strict Multi-tenant"
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def run_phase6_e2e():
    BASE_URL = "http://127.0.0.1:5173"
    print(f"=== Starting Phase 6 Document Extraction E2E Verification via Vite Proxy: {BASE_URL} ===")

    # 1. Register User A
    email_a = f"extract_user_{uuid.uuid4().hex[:6]}@cogniva.ai"
    password = "StrongPassword2026!"
    reg_body = json.dumps({"email": email_a, "password": password, "full_name": "Extraction Tester"}).encode("utf-8")
    req = urllib.request.Request(f"{BASE_URL}/api/v1/auth/register", data=reg_body, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req) as resp:
        user_a = json.loads(resp.read().decode())
        print(f"1. Registered User A: {user_a['email']} (id: {user_a['id']})")

    # 2. Login User A
    login_body = json.dumps({"email": email_a, "password": password}).encode("utf-8")
    req = urllib.request.Request(f"{BASE_URL}/api/v1/auth/login", data=login_body, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req) as resp:
        token_a = json.loads(resp.read().decode())["access_token"]
        headers_a = {"Authorization": f"Bearer {token_a}"}
        print("2. Logged in User A successfully.")

    # 3. Upload 2-page PDF
    pdf_bytes = create_sample_pdf()
    body, ct = encode_multipart_formdata({}, {"file": ("architecture_v2.pdf", pdf_bytes, "application/pdf")})
    req = urllib.request.Request(f"{BASE_URL}/api/v1/documents", data=body, headers={"Authorization": f"Bearer {token_a}", "Content-Type": ct})
    with urllib.request.urlopen(req) as resp:
        assert resp.status == 201
        pdf_doc = json.loads(resp.read().decode())
        assert pdf_doc["status"] == "uploaded"
        pdf_doc_id = pdf_doc["id"]
        print(f"3. Uploaded 2-page PDF: id={pdf_doc_id}, status={pdf_doc['status']}")

    # 4. Trigger PDF text extraction via POST /api/v1/documents/{id}/process
    req = urllib.request.Request(f"{BASE_URL}/api/v1/documents/{pdf_doc_id}/process", headers=headers_a, method="POST")
    with urllib.request.urlopen(req) as resp:
        assert resp.status == 200
        proc_doc = json.loads(resp.read().decode())
        assert proc_doc["status"] == "processed"
        assert proc_doc["metadata"]["total_pages"] == 2
        print(f"4. Processed PDF: status={proc_doc['status']}, pages={proc_doc['metadata']['total_pages']}")

    # 5. Fetch normalized extracted content for PDF
    req = urllib.request.Request(f"{BASE_URL}/api/v1/documents/{pdf_doc_id}/extracted", headers=headers_a)
    with urllib.request.urlopen(req) as resp:
        assert resp.status == 200
        ext_pdf = json.loads(resp.read().decode())
        assert ext_pdf["total_pages"] == 2
        assert len(ext_pdf["units"]) == 2
        # Check 1-indexed page preservation
        assert ext_pdf["units"][0]["page_number"] == 1
        assert "Page 1 Core Principles" in ext_pdf["units"][0]["text"]
        assert ext_pdf["units"][1]["page_number"] == 2
        assert "Page 2 Citation" in ext_pdf["units"][1]["text"]
        print("5. Verified PDF page-number preservation: Page 1 and Page 2 extracted cleanly.")

    # 6. Upload DOCX
    docx_bytes = create_sample_docx()
    body, ct = encode_multipart_formdata({}, {"file": ("knowledge_protocol.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    req = urllib.request.Request(f"{BASE_URL}/api/v1/documents", data=body, headers={"Authorization": f"Bearer {token_a}", "Content-Type": ct})
    with urllib.request.urlopen(req) as resp:
        assert resp.status == 201
        docx_doc = json.loads(resp.read().decode())
        docx_doc_id = docx_doc["id"]
        print(f"6. Uploaded DOCX: id={docx_doc_id}")

    # 7. Trigger DOCX text extraction via POST /api/v1/documents/{id}/process
    req = urllib.request.Request(f"{BASE_URL}/api/v1/documents/{docx_doc_id}/process", headers=headers_a, method="POST")
    with urllib.request.urlopen(req) as resp:
        assert resp.status == 200
        proc_docx = json.loads(resp.read().decode())
        assert proc_docx["status"] == "processed"
        assert proc_docx["metadata"]["total_pages"] is None
        print(f"7. Processed DOCX: status={proc_docx['status']}, total_units={proc_docx['metadata']['total_units']}")

    # 8. Fetch normalized extracted content for DOCX
    req = urllib.request.Request(f"{BASE_URL}/api/v1/documents/{docx_doc_id}/extracted", headers=headers_a)
    with urllib.request.urlopen(req) as resp:
        assert resp.status == 200
        ext_docx = json.loads(resp.read().decode())
        assert ext_docx["total_pages"] is None
        # Check units
        types = [u["unit_type"] for u in ext_docx["units"]]
        assert "heading" in types
        assert "paragraph" in types
        assert "table" in types
        # Check all page_numbers are None (never fabricated)
        assert all(u["page_number"] is None for u in ext_docx["units"])
        print("8. Verified DOCX structural preservation: Heading, Paragraph, Table extracted with page_number=None.")

    # 9. Multi-tenant isolation test: User B cannot process or view User A's extracted content
    email_b = f"intruder_{uuid.uuid4().hex[:6]}@cogniva.ai"
    reg_body = json.dumps({"email": email_b, "password": password, "full_name": "Intruder"}).encode("utf-8")
    req = urllib.request.Request(f"{BASE_URL}/api/v1/auth/register", data=reg_body, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req) as resp:
        pass
    login_body = json.dumps({"email": email_b, "password": password}).encode("utf-8")
    req = urllib.request.Request(f"{BASE_URL}/api/v1/auth/login", data=login_body, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req) as resp:
        token_b = json.loads(resp.read().decode())["access_token"]
        headers_b = {"Authorization": f"Bearer {token_b}"}

    # User B attempt to process User A's document
    req = urllib.request.Request(f"{BASE_URL}/api/v1/documents/{pdf_doc_id}/process", headers=headers_b, method="POST")
    try:
        urllib.request.urlopen(req)
        raise AssertionError("User B should not process User A's document")
    except urllib.error.HTTPError as err:
        assert err.code == 404, f"Expected 404, got {err.code}"
        print("9a. Cross-user processing attempt rejected with HTTP 404.")

    # User B attempt to view User A's extracted content
    req = urllib.request.Request(f"{BASE_URL}/api/v1/documents/{pdf_doc_id}/extracted", headers=headers_b)
    try:
        urllib.request.urlopen(req)
        raise AssertionError("User B should not view User A's extracted content")
    except urllib.error.HTTPError as err:
        assert err.code == 404, f"Expected 404, got {err.code}"
        print("9b. Cross-user extracted content view attempt rejected with HTTP 404.")

    print("\n>>> ALL PHASE 6 END-TO-END TESTS PASSED VIA VITE PROXY! <<<")

if __name__ == "__main__":
    run_phase6_e2e()
